from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "scripts" / "slide_assets_extended"
OUTPUT = ROOT / "VSP_EventOps_Presentation_Extended.pptx"


def ensure_assets():
    ASSETS.mkdir(parents=True, exist_ok=True)


def make_charts():
    # API checks
    fig, ax = plt.subplots(figsize=(8.5, 4.5))
    labels, values = ["Passed", "Failed"], [89, 0]
    bars = ax.bar(labels, values, color=["#16A34A", "#DC2626"], width=0.55)
    ax.set_ylim(0, 95)
    ax.set_title("API Regression & Stress Outcome", fontsize=16, weight="bold")
    ax.grid(axis="y", alpha=0.2)
    for b, v in zip(bars, values):
        ax.text(b.get_x() + b.get_width() / 2, v + 1, str(v), ha="center", va="bottom", fontsize=12)
    fig.tight_layout()
    fig.savefig(ASSETS / "api_pass_fail.png", dpi=220)
    plt.close(fig)

    # Dataset counts
    fig, ax = plt.subplots(figsize=(8.5, 4.8))
    cats = ["Students", "Events", "Staff", "Notifications"]
    vals = [407, 132, 19, 10]
    bars = ax.barh(cats, vals, color=["#1D4ED8", "#2563EB", "#60A5FA", "#93C5FD"])
    ax.set_xlim(0, 440)
    ax.set_title("Seeded Dataset Sizes", fontsize=16, weight="bold")
    ax.grid(axis="x", alpha=0.2)
    for b, v in zip(bars, vals):
        ax.text(v + 4, b.get_y() + b.get_height() / 2, str(v), va="center", fontsize=11)
    fig.tight_layout()
    fig.savefig(ASSETS / "dataset_sizes.png", dpi=220)
    plt.close(fig)

    # RBAC
    fig, ax = plt.subplots(figsize=(8.5, 4.5))
    labels = ["Total checks", "GEO forbidden", "Shared allowed", "Admin forbidden"]
    vals = [28, 15, 13, 0]
    bars = ax.bar(labels, vals, color=["#111827", "#DC2626", "#16A34A", "#6B7280"])
    ax.set_ylim(0, 31)
    ax.set_title("RBAC Validation Summary", fontsize=16, weight="bold")
    ax.grid(axis="y", alpha=0.2)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + 0.5, str(v), ha="center", va="bottom", fontsize=11)
    fig.tight_layout()
    fig.savefig(ASSETS / "rbac_summary.png", dpi=220)
    plt.close(fig)


def style_title(shape):
    p = shape.text_frame.paragraphs[0]
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(17, 24, 39)


def add_footer(slide):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(6.9), Inches(12.4), Inches(0.28))
    tf = box.text_frame
    tf.text = "EventOps | NYU Abu Dhabi | Rabeya Zahan Mily"
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(107, 114, 128)


def add_bullets(slide, items, left=0.9, top=1.8, width=11.4, height=4.8, size=22):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.add_paragraph() if i else tf.paragraphs[0]
        p.text = it
        p.font.size = Pt(size)
        p.level = 0


def add_title_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "EventOps"
    s.placeholders[1].text = (
        "A Full-Stack Operations Platform for the Visiting Students Program\n"
        "NYU Abu Dhabi | Capstone Presentation | Spring 2026"
    )
    style_title(s.shapes.title)
    for p in s.placeholders[1].text_frame.paragraphs:
        p.font.size = Pt(20)
    add_footer(s)


def add_section_header(prs, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    style_title(s.shapes.title)
    box = s.shapes.add_textbox(Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.4))
    tf = box.text_frame
    tf.text = subtitle
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = RGBColor(55, 65, 81)
    add_footer(s)


def add_two_col_slide(prs, title, left_title, left_points, right_title, right_points):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    style_title(s.shapes.title)
    # left card
    l = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.5), Inches(6.0), Inches(4.9))
    l.fill.solid(); l.fill.fore_color.rgb = RGBColor(239, 246, 255); l.line.color.rgb = RGBColor(147, 197, 253)
    lt = l.text_frame; lt.text = left_title
    lt.paragraphs[0].font.bold = True; lt.paragraphs[0].font.size = Pt(20)
    for pt in left_points:
        p = lt.add_paragraph(); p.text = pt; p.font.size = Pt(16)
    # right card
    r = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.95), Inches(1.5), Inches(5.7), Inches(4.9))
    r.fill.solid(); r.fill.fore_color.rgb = RGBColor(240, 253, 244); r.line.color.rgb = RGBColor(134, 239, 172)
    rt = r.text_frame; rt.text = right_title
    rt.paragraphs[0].font.bold = True; rt.paragraphs[0].font.size = Pt(20)
    for pt in right_points:
        p = rt.add_paragraph(); p.text = pt; p.font.size = Pt(16)
    add_footer(s)


def add_table_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Admin vs GEO Access Matrix"
    style_title(s.shapes.title)
    t = s.shapes.add_table(9, 3, Inches(0.7), Inches(1.5), Inches(12.0), Inches(4.95)).table
    t.columns[0].width = Inches(7.4)
    t.columns[1].width = Inches(2.2)
    t.columns[2].width = Inches(2.4)
    headers = ["Feature / Operation", "Admin", "GEO"]
    for c, h in enumerate(headers):
        cell = t.cell(0, c); cell.text = h
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(30, 58, 138)
        p = cell.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(15); p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER
    rows = [
        ("Students list/detail", "Yes", "Yes"),
        ("Create / delete student", "Yes", "No"),
        ("Events list/detail", "Yes", "Yes"),
        ("Create / lock / unlock event", "Yes", "No"),
        ("Bulk assign / clear assignments", "Yes", "No"),
        ("System settings + admin maintenance", "Yes", "No"),
        ("Attendance export", "Yes", "Yes"),
        ("Personal agenda", "Yes", "Yes"),
    ]
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
            p = t.cell(r, c).text_frame.paragraphs[0]
            p.font.size = Pt(14)
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
    add_footer(s)


def add_chart_slide(prs, title, img_name, note):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    style_title(s.shapes.title)
    s.shapes.add_picture(str(ASSETS / img_name), Inches(1.0), Inches(1.65), width=Inches(11.3))
    box = s.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.3), Inches(0.55))
    box.text_frame.text = note
    p = box.text_frame.paragraphs[0]
    p.font.size = Pt(16); p.alignment = PP_ALIGN.CENTER
    add_footer(s)


def add_timeline(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Project Timeline (January–April)"
    style_title(s.shapes.title)
    items = [
        ("January", "Requirements, problem definition, architecture, DB schema"),
        ("February", "Backend + frontend foundations, auth and role model"),
        ("March", "Attendance/strike workflows, modules, UX refinement"),
        ("April", "Stress tests, RBAC validation, report and final polish"),
    ]
    y = 1.75
    for m, d in items:
        chip = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(y), Inches(2.0), Inches(0.72))
        chip.fill.solid(); chip.fill.fore_color.rgb = RGBColor(30, 64, 175); chip.line.fill.background()
        chip.text = m
        cp = chip.text_frame.paragraphs[0]
        cp.font.size = Pt(15); cp.font.bold = True; cp.font.color.rgb = RGBColor(255, 255, 255); cp.alignment = PP_ALIGN.CENTER
        txt = s.shapes.add_textbox(Inches(3.25), Inches(y + 0.02), Inches(8.8), Inches(0.85))
        txt.text_frame.text = d
        txt.text_frame.paragraphs[0].font.size = Pt(18)
        y += 1.2
    add_footer(s)


def build():
    ensure_assets()
    make_charts()
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)
    add_section_header(prs, "Presentation Roadmap", "Problem → Solution → Architecture → Evaluation → Impact")
    add_bullets(
        prs.slides.add_slide(prs.slide_layouts[5]),
        [
            "Manual workflows were fragmented across spreadsheets and email.",
            "Data updates were duplicated and hard to reconcile.",
            "Staff lacked real-time visibility into operations.",
            "EventOps objective: one trusted platform for VSP operations.",
        ],
        size=21,
    )
    prs.slides[-1].shapes.title.text = "Problem & Motivation"
    style_title(prs.slides[-1].shapes.title)
    add_footer(prs.slides[-1])

    add_two_col_slide(
        prs,
        "Solution Overview",
        "What EventOps Delivers",
        [
            "Centralized student and cohort management",
            "Event planning, assignments, attendance workflows",
            "Strike processing and risk visibility",
            "Unified dashboard for daily operations",
        ],
        "Why It Matters",
        [
            "Reduces repetitive manual coordination",
            "Improves data consistency and trust",
            "Clarifies role ownership and accountability",
            "Creates a scalable operational foundation",
        ],
    )

    add_section_header(prs, "Technical Design", "Full-stack architecture built for maintainability and control")
    add_two_col_slide(
        prs,
        "Architecture Layers",
        "Frontend",
        [
            "Next.js + TypeScript",
            "Role-aware UX and dashboards",
            "Reusable components and state stores",
        ],
        "Backend + Data",
        [
            "Node/Express APIs with RBAC",
            "Session-based authentication",
            "PostgreSQL with Sequelize models",
        ],
    )
    add_table_slide(prs)

    add_section_header(prs, "Evaluation Results", "Evidence from stress testing, dataset validation, and RBAC checks")
    add_chart_slide(prs, "API Reliability Results", "api_pass_fail.png", "89 checks passed, 0 failed.")
    add_chart_slide(prs, "Dataset Scale Used in Testing", "dataset_sizes.png", "Seeded data mirrors real operational volume.")
    add_chart_slide(prs, "Access-Control Validation", "rbac_summary.png", "Admin-only paths blocked for GEO; shared paths available.")

    add_two_col_slide(
        prs,
        "Limitations & Challenges",
        "Current Constraints",
        [
            "Single-semester implementation window",
            "Evaluation in controlled local environment",
            "Long-term production observability pending",
        ],
        "How They Were Managed",
        [
            "Prioritized high-impact workflows first",
            "Used regression/stress tests to reduce risk",
            "Defined roadmap for future hardening",
        ],
    )
    add_timeline(prs)
    add_two_col_slide(
        prs,
        "Impact & Next Steps",
        "Current Impact",
        [
            "One operational source of truth",
            "Faster staff coordination and handoffs",
            "Safer role boundaries and auditability",
        ],
        "Future Extensions",
        [
            "Production monitoring and alerting",
            "Institutional SSO and wider integrations",
            "Advanced reporting and predictive analytics",
        ],
    )
    add_section_header(prs, "Thank You", "Questions?")

    prs.save(str(OUTPUT))
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build()

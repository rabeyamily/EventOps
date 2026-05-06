from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "scripts" / "slide_assets"
OUTPUT = ROOT / "VSP_EventOps_Presentation.pptx"
PDF_PREVIEW = ROOT / "scripts" / "report-page.png"


def ensure_assets():
    ASSETS.mkdir(parents=True, exist_ok=True)


def create_chart_images():
    # API stress chart
    fig, ax = plt.subplots(figsize=(8, 4.2))
    labels = ["Passed", "Failed"]
    values = [89, 0]
    colors = ["#1E8E3E", "#D93025"]
    bars = ax.bar(labels, values, color=colors, width=0.6)
    ax.set_title("API Stress Summary", fontsize=16, weight="bold")
    ax.set_ylabel("Checks")
    ax.set_ylim(0, 95)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 1, str(value), ha="center", va="bottom", fontsize=12)
    ax.grid(axis="y", alpha=0.2)
    fig.tight_layout()
    fig.savefig(ASSETS / "api_stress_chart.png", dpi=220)
    plt.close(fig)

    # Dataset counts chart
    fig, ax = plt.subplots(figsize=(8, 4.6))
    entities = ["Staff", "Students", "Events", "Notifications"]
    counts = [19, 407, 132, 10]
    bars = ax.barh(entities, counts, color=["#1A73E8", "#174EA6", "#5F6368", "#1E8E3E"])
    ax.set_title("Post-import Dataset Counts", fontsize=16, weight="bold")
    for bar, value in zip(bars, counts):
        ax.text(value + 4, bar.get_y() + bar.get_height() / 2, str(value), va="center", fontsize=11)
    ax.set_xlim(0, 440)
    ax.grid(axis="x", alpha=0.2)
    fig.tight_layout()
    fig.savefig(ASSETS / "dataset_counts_chart.png", dpi=220)
    plt.close(fig)

    # Role access chart
    fig, ax = plt.subplots(figsize=(8, 4.2))
    labels = ["Total Checks", "GEO Forbidden", "Shared Allowed", "Admin Forbidden"]
    values = [28, 15, 13, 0]
    colors = ["#1A73E8", "#D93025", "#1E8E3E", "#5F6368"]
    bars = ax.bar(labels, values, color=colors)
    ax.set_title("RBAC Validation Summary", fontsize=16, weight="bold")
    ax.set_ylim(0, 31)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 0.6, str(value), ha="center", va="bottom", fontsize=11)
    ax.grid(axis="y", alpha=0.2)
    fig.tight_layout()
    fig.savefig(ASSETS / "rbac_chart.png", dpi=220)
    plt.close(fig)


def style_title(title_shape):
    p = title_shape.text_frame.paragraphs[0]
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = RGBColor(17, 24, 39)


def add_footer(slide, text="EventOps | NYU Abu Dhabi | Spring 2026"):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(6.85), Inches(12.5), Inches(0.3))
    tf = box.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(107, 114, 128)
    p.alignment = PP_ALIGN.RIGHT


def add_agenda_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Agenda"
    style_title(slide.shapes.title)
    agenda = [
        "1. Problem & Motivation",
        "2. System Architecture",
        "3. Core Features & Access Control",
        "4. Database Model",
        "5. Evaluation Results",
        "6. Limitations, Impact, Next Steps",
    ]
    box = slide.shapes.add_textbox(Inches(0.9), Inches(1.7), Inches(11.4), Inches(4.8))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(agenda):
        p = tf.add_paragraph() if i else tf.paragraphs[0]
        p.text = item
        p.font.size = Pt(28)
        p.font.color.rgb = RGBColor(31, 41, 55)
        p.level = 0
    add_footer(slide)


def add_problem_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Problem & Motivation"
    style_title(slide.shapes.title)
    points = [
        "VSP operations relied on spreadsheets, emails, and manual coordination.",
        "Data duplication and inconsistent updates created operational risk.",
        "Staff lacked a single source of truth across student/event workflows.",
        "Goal: build a practical staff-first platform to centralize operations.",
    ]
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(7.2), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(points):
        p = tf.add_paragraph() if i else tf.paragraphs[0]
        p.text = item
        p.font.size = Pt(22)
        p.level = 0
    if PDF_PREVIEW.exists():
        slide.shapes.add_picture(str(PDF_PREVIEW), Inches(8.2), Inches(1.7), width=Inches(4.4))
    add_footer(slide)


def add_architecture_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "System Architecture"
    style_title(slide.shapes.title)

    layers = [
        ("Frontend (Next.js + TypeScript)", RGBColor(219, 234, 254), Inches(0.9)),
        ("Backend API (Node/Express + RBAC)", RGBColor(220, 252, 231), Inches(2.35)),
        ("PostgreSQL + Sequelize Models", RGBColor(254, 243, 199), Inches(3.8)),
    ]
    for text, color, top in layers:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.0), top, Inches(10.8), Inches(1.1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = RGBColor(75, 85, 99)
        tf = shape.text_frame
        tf.text = text
        p = tf.paragraphs[0]
        p.font.size = Pt(24)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

    box = slide.shapes.add_textbox(Inches(0.95), Inches(5.2), Inches(11), Inches(1.25))
    tf = box.text_frame
    tf.text = "Three-layer architecture improves maintainability, testing, and scalability."
    p = tf.paragraphs[0]
    p.font.size = Pt(20)
    p.font.color.rgb = RGBColor(31, 41, 55)
    p.alignment = PP_ALIGN.CENTER
    add_footer(slide)


def add_features_table_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Feature Modules & Role Access"
    style_title(slide.shapes.title)

    rows = 8
    cols = 3
    table = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(1.6), Inches(12.1), Inches(4.8)).table
    table.columns[0].width = Inches(7.5)
    table.columns[1].width = Inches(2.2)
    table.columns[2].width = Inches(2.4)

    headers = ["Feature / Operation", "Admin", "GEO"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 64, 175)
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

    data = [
        ("Students list/detail", "Yes", "Yes"),
        ("Create/Delete student", "Yes", "No"),
        ("Events list/detail", "Yes", "Yes"),
        ("Create/Lock/Unlock event", "Yes", "No"),
        ("Attendance export", "Yes", "Yes"),
        ("System settings + Admin maintenance", "Yes", "No"),
        ("Personal agenda", "Yes", "Yes"),
    ]
    for r, row in enumerate(data, start=1):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
            p = table.cell(r, c).text_frame.paragraphs[0]
            p.font.size = Pt(14)
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
    add_footer(slide)


def add_data_model_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Database Model (Core Entities)"
    style_title(slide.shapes.title)

    entities = [
        ("Staff", 0.8, 2.0),
        ("Students", 3.2, 2.0),
        ("Events", 5.6, 2.0),
        ("Attendance Sheets", 8.0, 2.0),
        ("Attendances", 10.0, 3.5),
        ("Strikes", 5.6, 4.6),
    ]
    for name, left, top in entities:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(2.0), Inches(0.9))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(243, 244, 246)
        shape.line.color.rgb = RGBColor(55, 65, 81)
        tf = shape.text_frame
        tf.text = name
        p = tf.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(14)
        p.alignment = PP_ALIGN.CENTER

    box = slide.shapes.add_textbox(Inches(0.8), Inches(5.9), Inches(11.5), Inches(1.0))
    tf = box.text_frame
    tf.text = "Key relations: Events↔Attendances, Students↔Attendances, Events↔Strikes, Staff↔all operational records."
    p = tf.paragraphs[0]
    p.font.size = Pt(17)
    add_footer(slide)


def add_chart_slide(prs, title, image_name, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    style_title(slide.shapes.title)
    slide.shapes.add_picture(str(ASSETS / image_name), Inches(1.0), Inches(1.6), width=Inches(11.3))
    box = slide.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.2), Inches(0.6))
    tf = box.text_frame
    tf.text = subtitle
    p = tf.paragraphs[0]
    p.font.size = Pt(16)
    p.alignment = PP_ALIGN.CENTER
    add_footer(slide)


def add_timeline_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Project Timeline (Jan–Apr 2026)"
    style_title(slide.shapes.title)

    entries = [
        ("January", "Scope, requirements, architecture, schema design"),
        ("February", "Backend + frontend foundations, auth/RBAC, core APIs"),
        ("March", "Attendance/strike workflows, modules, UI refinement"),
        ("April", "Stress tests, RBAC validation, report finalization"),
    ]
    y = 1.7
    for month, desc in entries:
        chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(y), Inches(2.0), Inches(0.75))
        chip.fill.solid()
        chip.fill.fore_color.rgb = RGBColor(30, 64, 175)
        chip.line.fill.background()
        chip.text = month
        chip.text_frame.paragraphs[0].font.size = Pt(16)
        chip.text_frame.paragraphs[0].font.bold = True
        chip.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        chip.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        text = slide.shapes.add_textbox(Inches(3.2), Inches(y + 0.03), Inches(8.7), Inches(0.9))
        text.text_frame.text = desc
        text.text_frame.paragraphs[0].font.size = Pt(18)
        y += 1.25
    add_footer(slide)


def add_closing_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Conclusion & Next Steps"
    style_title(slide.shapes.title)
    points = [
        "EventOps successfully replaces fragmented admin workflows with one platform.",
        "Testing validated reliability (89/89), RBAC boundaries, and operational readiness.",
        "Future work: production observability, SSO integration, deeper analytics.",
        "Outcome: higher visibility, consistency, and coordination for VSP staff.",
    ]
    box = slide.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(11.2), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    for i, pt in enumerate(points):
        p = tf.add_paragraph() if i else tf.paragraphs[0]
        p.text = pt
        p.font.size = Pt(23)
        p.level = 0
    thanks = slide.shapes.add_textbox(Inches(0.9), Inches(6.1), Inches(11.2), Inches(0.7))
    thanks.text_frame.text = "Thank you"
    thanks.text_frame.paragraphs[0].font.size = Pt(30)
    thanks.text_frame.paragraphs[0].font.bold = True
    thanks.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_footer(slide)


def build_presentation():
    ensure_assets()
    create_chart_images()

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "EventOps"
    slide.placeholders[1].text = (
        "A Full-Stack Operations Platform for the Visiting Students Program\n"
        "NYU Abu Dhabi | Rabeya Zahan Mily | Spring 2026"
    )
    style_title(slide.shapes.title)
    for p in slide.placeholders[1].text_frame.paragraphs:
        p.font.size = Pt(22)
    add_footer(slide)

    add_agenda_slide(prs)
    add_problem_slide(prs)
    add_architecture_slide(prs)
    add_features_table_slide(prs)
    add_data_model_slide(prs)
    add_chart_slide(
        prs,
        "Evaluation: API Reliability",
        "api_stress_chart.png",
        "All 89 API checks passed with zero failures.",
    )
    add_chart_slide(
        prs,
        "Evaluation: Dataset Scale",
        "dataset_counts_chart.png",
        "Post-import data confirms realistic semester-scale testing.",
    )
    add_chart_slide(
        prs,
        "Evaluation: Access Control",
        "rbac_chart.png",
        "RBAC checks confirm admin-only actions are blocked for GEO users.",
    )
    add_timeline_slide(prs)
    add_closing_slide(prs)

    prs.save(str(OUTPUT))
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_presentation()
